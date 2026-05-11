import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from config import (
    COMPLIANCE_LEDGER_EXCEL_PATH,
    COMPLIANCE_LEDGER_JSON_PATH,
    COMPLIANCE_RESPONSIBLE_PERSONS_PATH,
    MODEL_CHAT,
)
from file_store import atomic_write_text, file_lock


DEFAULT_RESPONSIBLE_PERSONS = {
    "规划与资产部/深化改革领导小组办公室": "富小鹏",
    "财务部": "杨焕",
    "审计部/法务合规部": "李莹",
    "人力资源部": "陈锐",
    "党群办公室/董事会办公室/行政办公室": "刘芳",
    "安全质量部": "霍晓冬",
    "纪委办公室/巡察工作领导小组办公室": "边宁",
}

VALID_IMPLEMENTATIONS = {"已按要求补充完善", "未见落实", "不涉及", "/"}
WORKSHEET_NAME = "合规管理牵头部门合规审查台账"
HEADERS = ["序号", "重大事项", "程序", "审查时间", "审查单位", "合规审查意见", "具体意见描述", "落实情况", "承办单位", "背景材料"]


def load_responsible_persons(path: str | Path = COMPLIANCE_RESPONSIBLE_PERSONS_PATH) -> dict[str, str]:
    p = Path(path)
    if not p.exists():
        return dict(DEFAULT_RESPONSIBLE_PERSONS)
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return dict(DEFAULT_RESPONSIBLE_PERSONS)
    if not isinstance(data, dict):
        return dict(DEFAULT_RESPONSIBLE_PERSONS)
    result = dict(DEFAULT_RESPONSIBLE_PERSONS)
    for dept, person in data.items():
        dept_text = str(dept).strip()
        person_text = str(person).strip()
        if dept_text and person_text:
            result[dept_text] = person_text
    return result


def save_responsible_persons(persons: dict[str, str], path: str | Path = COMPLIANCE_RESPONSIBLE_PERSONS_PATH) -> dict[str, str]:
    cleaned = {}
    for dept, person in persons.items():
        dept_text = str(dept).strip()
        person_text = str(person).strip()
        if dept_text and person_text:
            cleaned[dept_text] = person_text
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with file_lock(p):
        atomic_write_text(p, json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")
    return cleaned


def normalize_review_opinion(opinion_text: str | None) -> str:
    text = re.sub(r"\s+", "", opinion_text or "")
    if any(word in text for word in ["不同意", "不予同意", "暂不同意"]):
        return "不予同意"
    if any(word in text for word in ["建议", "补充", "完善", "修改", "调整", "需进一步", "请进一步"]):
        return "建议补充完善"
    return "同意"


def _clean_text(value: Any, default: str = "") -> str:
    text = str(value or "").strip()
    return text or default


def _normalize_implementation(value: Any, opinion_text: str) -> str:
    text = _clean_text(value, "")
    if text in VALID_IMPLEMENTATIONS:
        return text
    if normalize_review_opinion(opinion_text) == "建议补充完善":
        return "未见落实"
    return "/"


def _detail_for_opinion(item: dict[str, Any]) -> str:
    detail = _clean_text(item.get("detail"), "")
    if detail:
        return detail
    opinion_text = _clean_text(item.get("opinion_text"), "")
    if normalize_review_opinion(opinion_text) == "同意":
        return "/"
    return opinion_text or "/"


def _row_from(role: str, item: dict[str, Any], department: str | None = None) -> dict[str, str]:
    dept = _clean_text(department if department is not None else item.get("department"), "")
    if role == "首席合规官":
        review_unit = "首席合规官"
    else:
        review_unit = f"{role}（{dept}）" if dept else role
    opinion_text = _clean_text(item.get("opinion_text"), "")
    return {
        "review_time": _clean_text(item.get("time"), ""),
        "review_unit": review_unit,
        "review_opinion": normalize_review_opinion(opinion_text),
        "detail": _detail_for_opinion(item),
        "implementation": _normalize_implementation(item.get("implementation"), opinion_text),
    }


def build_review_rows(item: dict[str, Any]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    chief = item.get("chief") or {}
    compliance = item.get("compliance") or {}
    countersign = item.get("countersign") or []
    undertaking = item.get("undertaking") or {}

    if chief:
        rows.append(_row_from("首席合规官", chief))
    if compliance:
        rows.append(_row_from("合规管理牵头部门", compliance, compliance.get("department") or "审计部/法务合规部"))
    for entry in countersign:
        if isinstance(entry, dict):
            rows.append(_row_from("会签单位", entry))
    if undertaking:
        rows.append(_row_from("承办单位", undertaking))

    return rows or [_row_from("承办单位", {})]


def _normalize_procedure(value: Any) -> str:
    text = _clean_text(value, "")
    if "董事" in text:
        return "董事会审议"
    return "总办会审议"


def normalize_extracted_item(raw: dict[str, Any]) -> dict[str, Any]:
    background = raw.get("background_materials") or raw.get("attachments") or []
    if isinstance(background, str):
        background_items = [background]
    else:
        background_items = [str(x).strip() for x in background if str(x).strip()]
    return {
        "title": _clean_text(raw.get("title") or raw.get("重大事项"), "未识别标题"),
        "procedure": _normalize_procedure(raw.get("procedure") or raw.get("程序")),
        "undertaking_department": "法务合规部",
        "background_materials": [re.sub(r"\.(pdf|docx?|xlsx?|xls)$", "", x, flags=re.IGNORECASE) for x in background_items],
        "review_rows": raw.get("review_rows") or build_review_rows(raw),
        "warnings": raw.get("warnings") or [],
    }


def _parse_json_object(raw: str) -> dict[str, Any]:
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("LLM 未返回 JSON 对象")
    return json.loads(match.group())


def _needs_ocr(text: str) -> bool:
    if not text.strip():
        return True
    chinese_count = len(re.findall(r"[\u4e00-\u9fff]", text))
    slash_tokens = len(re.findall(r"/\d+", text))
    return chinese_count < 20 or slash_tokens > chinese_count


def extract_pdf_text(pdf_bytes: bytes, filename: str, vision_model: str | None = None) -> str:
    from ledger_helpers import extract_file_text, ocr_pdf_with_vision

    text = extract_file_text(pdf_bytes, filename)
    if _needs_ocr(text):
        text = ocr_pdf_with_vision(pdf_bytes, model=vision_model)
    return text


def extract_compliance_item(text: str, responsible_persons: dict[str, str] | None = None) -> dict[str, Any]:
    from llm_client import get_llm_client

    persons = responsible_persons or load_responsible_persons()
    prompt = f"""你是企业合规审查台账录入助手。请从 OA 流程表单/审批记录中提取合规审查工作台账字段。

部门负责人配置：
{json.dumps(persons, ensure_ascii=False, indent=2)}

抽取规则：
1. title 填文件信息中的标题内容。
2. procedure 只能填“董事会审议”或“总办会审议”，根据正文判断。
3. undertaking 取“拟稿单位意见”中对应部门负责人的意见。
4. countersign 取会签意见中除“审计部/法务合规部”以外的部门负责人意见；多个部门逐个返回。
5. compliance 取会签意见中的“审计部/法务合规部”负责人意见。
6. chief 取胡鹏斌意见。
7. 每个意见对象包含 department、person、time、opinion_text、detail、implementation。
8. implementation 只能填“/”“已按要求补充完善”“未见落实”“不涉及”。
9. attachments 提取正文附件列表中的附件名称，去掉 PDF/DOC/XLS 等后缀。

只返回 JSON 对象，格式如下：
{{
  "title": "",
  "procedure": "董事会审议",
  "attachments": [],
  "undertaking": {{"department": "", "person": "", "time": "", "opinion_text": "", "detail": "", "implementation": "/"}},
  "countersign": [],
  "compliance": {{"department": "审计部/法务合规部", "person": "", "time": "", "opinion_text": "", "detail": "", "implementation": "/"}},
  "chief": {{"person": "胡鹏斌", "time": "", "opinion_text": "", "detail": "", "implementation": "/"}},
  "warnings": []
}}

PDF/OA内容：
{text[:16000]}"""
    client = get_llm_client()
    response = client.chat.completions.create(
        model=MODEL_CHAT,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
        max_tokens=3000,
    )
    raw = response.choices[0].message.content or ""
    return normalize_extracted_item(_parse_json_object(raw))


def load_records(path: str | Path = COMPLIANCE_LEDGER_JSON_PATH) -> list[dict[str, Any]]:
    p = Path(path)
    if not p.exists():
        return []
    try:
        with file_lock(p):
            data = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []
    return data if isinstance(data, list) else []


def save_records(records: list[dict[str, Any]], path: str | Path = COMPLIANCE_LEDGER_JSON_PATH) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with file_lock(p):
        atomic_write_text(p, json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def append_record(record: dict[str, Any], path: str | Path = COMPLIANCE_LEDGER_JSON_PATH) -> list[dict[str, Any]]:
    records = load_records(path)
    next_record = dict(record)
    next_record["sequence"] = len(records) + 1
    records.append(next_record)
    save_records(records, path)
    return records


def create_compliance_workbook(records: list[dict[str, Any]], output_path: str | Path = COMPLIANCE_LEDGER_EXCEL_PATH) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = WORKSHEET_NAME

    ws.merge_cells("A1:J1")
    ws["A1"] = "合规管理牵头部门合规审查工作台账"
    ws["A1"].font = Font(name="宋体", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 46.95

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(2, col, header)
        cell.font = Font(name="宋体", bold=True, size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[2].height = 40.8

    widths = [13, 28, 16, 14, 28, 19, 32, 20, 16, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

    row_idx = 3
    for idx, record in enumerate(records, start=1):
        review_rows = record.get("review_rows") or []
        if not review_rows:
            review_rows = [{"review_time": "", "review_unit": "", "review_opinion": "同意", "detail": "/", "implementation": "/"}]
        start = row_idx
        end = row_idx + len(review_rows) - 1

        merged_values = {
            1: record.get("sequence") or idx,
            2: record.get("title") or "",
            3: _normalize_procedure(record.get("procedure")),
            9: record.get("undertaking_department") or "法务合规部",
            10: "、".join(record.get("background_materials") or []) or "/",
        }
        for col, value in merged_values.items():
            ws.cell(start, col, value)
            if end > start:
                ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)

        for offset, review in enumerate(review_rows):
            r = start + offset
            ws.cell(r, 4, review.get("review_time") or "")
            ws.cell(r, 5, review.get("review_unit") or "")
            ws.cell(r, 6, review.get("review_opinion") or "同意")
            ws.cell(r, 7, review.get("detail") or "/")
            ws.cell(r, 8, review.get("implementation") or "/")
            ws.row_dimensions[r].height = 36

        for r in range(start, end + 1):
            for c in range(1, 11):
                cell = ws.cell(r, c)
                cell.font = Font(name="宋体", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
        row_idx = end + 1

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
